"""Geração de relatórios exportáveis (SPEC.md seção 14).

Três saídas, todas a partir dos mesmos dados coletados em uma janela de tempo:

- `build_incidents_xlsx`  — alertas do período (fonte, severidade, diagnóstico)
- `build_actions_xlsx`    — ações executadas, quem aprovou, resultado
- `build_summary_pdf`     — resumo de uma página para leitura/arquivamento

Os arquivos são gerados em memória (`BytesIO`) e devolvidos como bytes, de
modo que a API pode transmiti-los sem escrever em disco no servidor.

Nota de privacidade: os relatórios trazem `parameters` das ações, que podem
conter nomes de host e de banco. Não contêm segredos — a camada de conectores
nunca coloca credenciais em `parameters` (SPEC 10.8) — mas trate o arquivo
exportado com o mesmo cuidado de qualquer documento interno de infraestrutura.
"""
from __future__ import annotations

import dataclasses
import io
from datetime import datetime, timedelta, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from sai.db.models import Action, Alert, AuditLog

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclasses.dataclass(frozen=True, slots=True)
class ReportPeriod:
    """Janela de tempo do relatório, sempre em UTC."""

    start: datetime
    end: datetime

    @classmethod
    def last_days(cls, days: int = 30) -> "ReportPeriod":
        end = datetime.now(timezone.utc)
        return cls(start=end - timedelta(days=days), end=end)

    @classmethod
    def month(cls, year: int, month: int) -> "ReportPeriod":
        start = datetime(year, month, 1, tzinfo=timezone.utc)
        end = datetime(year + (month == 12), (month % 12) + 1, 1, tzinfo=timezone.utc)
        return cls(start=start, end=end)

    @property
    def label(self) -> str:
        return f"{self.start.strftime('%d/%m/%Y')} a {self.end.strftime('%d/%m/%Y')}"


@dataclasses.dataclass(slots=True)
class ReportData:
    period: ReportPeriod
    alerts: list[Alert]
    actions: list[Action]
    audit_entries: list[AuditLog]

    @property
    def executed_actions(self) -> list[Action]:
        return [a for a in self.actions if a.status in ("succeeded", "failed")]

    @property
    def stats(self) -> dict[str, int]:
        return {
            "alertas": len(self.alerts),
            "alertas_criticos": sum(1 for a in self.alerts if a.severity == "critical"),
            "acoes_propostas": len(self.actions),
            "acoes_executadas": len(self.executed_actions),
            "acoes_com_sucesso": sum(1 for a in self.actions if a.status == "succeeded"),
            "acoes_com_falha": sum(1 for a in self.actions if a.status == "failed"),
            "acoes_rejeitadas": sum(1 for a in self.actions if a.status == "rejected"),
        }


async def collect_report_data(session: AsyncSession, period: ReportPeriod) -> ReportData:
    """Lê alertas, ações e auditoria da janela informada."""
    alerts = (
        await session.execute(
            select(Alert)
            .where(Alert.created_at >= period.start, Alert.created_at < period.end)
            .order_by(Alert.created_at)
        )
    ).scalars().all()

    actions = (
        await session.execute(
            select(Action)
            .where(Action.created_at >= period.start, Action.created_at < period.end)
            .order_by(Action.created_at)
        )
    ).scalars().all()

    audit_entries = (
        await session.execute(
            select(AuditLog)
            .where(AuditLog.created_at >= period.start, AuditLog.created_at < period.end)
            .order_by(AuditLog.created_at)
        )
    ).scalars().all()

    return ReportData(
        period=period,
        alerts=list(alerts),
        actions=list(actions),
        audit_entries=list(audit_entries),
    )


# ---------------------------------------------------------------- Excel


def _write_sheet(ws: Any, headers: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"


def _fmt(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y %H:%M") if value else ""


def build_incidents_xlsx(data: ReportData) -> bytes:
    """Planilha de incidentes (alertas) do período."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Incidentes"

    _write_sheet(
        ws,
        ["Data", "Fonte", "Severidade", "Resumo", "Status", "Diagnóstico"],
        [
            [
                _fmt(a.created_at),
                a.source,
                a.severity,
                a.summary,
                a.status,
                (a.diagnosis or "")[:1000],
            ]
            for a in data.alerts
        ],
        widths=[18, 14, 12, 50, 18, 80],
    )
    return _workbook_bytes(wb)


def build_actions_xlsx(data: ReportData) -> bytes:
    """Planilha de ações: o que foi proposto, quem aprovou, o que aconteceu."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Ações"

    _write_sheet(
        ws,
        [
            "Data",
            "Ferramenta",
            "Risco",
            "Descrição",
            "Exigiu aprovação",
            "Aprovado por",
            "Aprovado em",
            "Executado em",
            "Status",
            "Parâmetros",
            "Resultado",
        ],
        [
            [
                _fmt(a.created_at),
                a.tool_name,
                a.risk_level,
                a.proposed_description,
                "Sim" if a.requires_approval else "Não (allowlist)",
                str(a.approved_by) if a.approved_by else "",
                _fmt(a.approved_at),
                _fmt(a.executed_at),
                a.status,
                str(a.parameters),
                str(a.execution_result or "")[:1000],
            ]
            for a in data.actions
        ],
        widths=[18, 28, 10, 45, 18, 38, 18, 18, 14, 40, 50],
    )

    ws_audit = wb.create_sheet("Auditoria")
    _write_sheet(
        ws_audit,
        ["Data", "Ator", "Evento", "Entidade", "ID", "Payload"],
        [
            [
                _fmt(e.created_at),
                e.actor,
                e.event_type,
                e.entity_type,
                str(e.entity_id) if e.entity_id else "",
                str(e.payload)[:1000],
            ]
            for e in data.audit_entries
        ],
        widths=[18, 30, 28, 14, 38, 70],
    )
    return _workbook_bytes(wb)


def _workbook_bytes(wb: Workbook) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------- PDF


def build_summary_pdf(data: ReportData) -> bytes:
    """Resumo executivo de uma página: números do período, alertas críticos e
    as ações de maior risco executadas."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
        title=f"SAI — Relatório {data.period.label}",
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("SaiTitle", parent=styles["Title"], fontSize=18, spaceAfter=4)
    small = ParagraphStyle("SaiSmall", parent=styles["Normal"], fontSize=9, textColor=colors.grey)

    story: list[Any] = [
        Paragraph("SAI — Relatório de Operação", title_style),
        Paragraph(f"Período: {data.period.label}", small),
        Spacer(1, 0.8 * cm),
    ]

    stats = data.stats
    story.append(Paragraph("Resumo do período", styles["Heading2"]))
    story.append(
        _table(
            [["Indicador", "Total"]]
            + [[k.replace("_", " ").capitalize(), str(v)] for k, v in stats.items()],
            col_widths=[10 * cm, 4 * cm],
        )
    )
    story.append(Spacer(1, 0.7 * cm))

    critical = [a for a in data.alerts if a.severity == "critical"]
    story.append(Paragraph(f"Alertas críticos ({len(critical)})", styles["Heading2"]))
    if critical:
        story.append(
            _table(
                [["Data", "Fonte", "Resumo", "Status"]]
                + [
                    [_fmt(a.created_at), a.source, _clip(a.summary, 60), a.status]
                    for a in critical[:20]
                ],
                col_widths=[3 * cm, 2.5 * cm, 8 * cm, 3 * cm],
            )
        )
    else:
        story.append(Paragraph("Nenhum alerta crítico no período.", styles["Normal"]))
    story.append(Spacer(1, 0.7 * cm))

    high_risk = [a for a in data.executed_actions if a.risk_level in ("high", "critical")]
    story.append(Paragraph(f"Ações de alto risco executadas ({len(high_risk)})", styles["Heading2"]))
    if high_risk:
        story.append(
            _table(
                [["Data", "Ferramenta", "Risco", "Status"]]
                + [
                    [_fmt(a.executed_at), _clip(a.tool_name, 32), a.risk_level, a.status]
                    for a in high_risk[:20]
                ],
                col_widths=[3 * cm, 7 * cm, 2.5 * cm, 4 * cm],
            )
        )
    else:
        story.append(
            Paragraph("Nenhuma ação de risco alto ou crítico foi executada no período.", styles["Normal"])
        )

    story.append(Spacer(1, 1 * cm))
    story.append(
        Paragraph(
            f"Gerado automaticamente pelo SAI em {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC. "
            "Todas as ações listadas constam do log de auditoria imutável.",
            small,
        )
    )

    doc.build(story)
    return buffer.getvalue()


def _clip(text: str, limit: int) -> str:
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _table(rows: list[list[str]], col_widths: list[float]) -> Table:
    table = Table(rows, colWidths=col_widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3864")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B4C7E7")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F5FB")]),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    return table
