"""Testes da exportação de relatórios (SPEC 14).

Verificam que os arquivos gerados são realmente arquivos válidos daqueles
formatos — não apenas que a função não levantou exceção — e que os dados
chegam ao conteúdo.
"""
from __future__ import annotations

import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from sai.db.models import Action, Alert, AuditLog
from sai.reports import ReportPeriod, build_actions_xlsx, build_incidents_xlsx, build_summary_pdf
from sai.reports.builder import ReportData

NOW = datetime(2026, 9, 1, 12, 0, tzinfo=timezone.utc)


def _alert(summary: str, severity: str = "critical", diagnosis: str | None = "Causa raiz X") -> Alert:
    return Alert(
        id=uuid.uuid4(),
        source="zabbix",
        severity=severity,
        summary=summary,
        raw_payload={"_resource": "host-01"},
        diagnosis=diagnosis,
        status="open",
        created_at=NOW,
    )


def _action(tool: str, risk: str, status: str) -> Action:
    return Action(
        id=uuid.uuid4(),
        tool_name=tool,
        risk_level=risk,
        parameters={"host": "web-01"},
        proposed_description=f"Executar {tool}",
        requires_approval=True,
        status=status,
        created_at=NOW,
        executed_at=NOW + timedelta(minutes=2),
        execution_result={"ok": status == "succeeded"},
    )


@pytest.fixture
def report_data() -> ReportData:
    return ReportData(
        period=ReportPeriod(start=NOW - timedelta(days=30), end=NOW + timedelta(days=1)),
        alerts=[
            _alert("CPU acima de 95% em web-01"),
            _alert("Disco em 80%", severity="warning"),
        ],
        actions=[
            _action("linux_restart_service", "low", "succeeded"),
            _action("sql_restore_database", "critical", "succeeded"),
            _action("f5_disable_pool_member", "high", "failed"),
            _action("azure_resize_disk", "high", "rejected"),
        ],
        audit_entries=[
            AuditLog(
                id=1,
                actor="user:abc",
                event_type="action_approved",
                entity_type="action",
                entity_id=uuid.uuid4(),
                payload={"tool_name": "sql_restore_database"},
                created_at=NOW,
            )
        ],
    )


# -- estatísticas ---------------------------------------------------------


def test_stats_count_correctly(report_data: ReportData):
    stats = report_data.stats
    assert stats["alertas"] == 2
    assert stats["alertas_criticos"] == 1
    assert stats["acoes_propostas"] == 4
    assert stats["acoes_executadas"] == 3  # succeeded + failed, não a rejeitada
    assert stats["acoes_com_sucesso"] == 2
    assert stats["acoes_com_falha"] == 1
    assert stats["acoes_rejeitadas"] == 1


# -- Excel ----------------------------------------------------------------


def test_incidents_xlsx_is_a_valid_workbook_with_the_data(report_data: ReportData):
    content = build_incidents_xlsx(report_data)
    wb = load_workbook(io.BytesIO(content))
    ws = wb["Incidentes"]

    assert ws.max_row == 3  # cabeçalho + 2 alertas
    header = [c.value for c in ws[1]]
    assert header == ["Data", "Fonte", "Severidade", "Resumo", "Status", "Diagnóstico"]

    summaries = [ws.cell(row=r, column=4).value for r in range(2, ws.max_row + 1)]
    assert "CPU acima de 95% em web-01" in summaries


def test_actions_xlsx_has_both_sheets_and_audit_trail(report_data: ReportData):
    content = build_actions_xlsx(report_data)
    wb = load_workbook(io.BytesIO(content))

    assert wb.sheetnames == ["Ações", "Auditoria"]
    ws = wb["Ações"]
    assert ws.max_row == 5  # cabeçalho + 4 ações

    tools = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert "sql_restore_database" in tools

    audit = wb["Auditoria"]
    assert audit.max_row == 2
    assert audit.cell(row=2, column=3).value == "action_approved"


def test_allowlisted_action_is_labelled_as_such(report_data: ReportData):
    """Uma ação que rodou pela allowlist precisa ficar visualmente distinta de
    uma aprovada por uma pessoa — é a diferença entre 'alguém decidiu' e
    'a política decidiu'."""
    auto = _action("zabbix_acknowledge_problem", "low", "succeeded")
    auto.requires_approval = False
    report_data.actions.append(auto)

    wb = load_workbook(io.BytesIO(build_actions_xlsx(report_data)))
    ws = wb["Ações"]
    labels = [ws.cell(row=r, column=5).value for r in range(2, ws.max_row + 1)]
    assert "Não (allowlist)" in labels
    assert "Sim" in labels


def test_empty_period_still_produces_a_valid_file():
    """Um mês sem incidentes não pode gerar arquivo corrompido."""
    empty = ReportData(
        period=ReportPeriod.last_days(30), alerts=[], actions=[], audit_entries=[]
    )
    wb = load_workbook(io.BytesIO(build_incidents_xlsx(empty)))
    assert wb["Incidentes"].max_row == 1  # só o cabeçalho


# -- PDF ------------------------------------------------------------------


def test_summary_pdf_is_a_valid_pdf(report_data: ReportData):
    content = build_summary_pdf(report_data)
    assert content.startswith(b"%PDF-"), "saída não é um PDF"
    assert b"%%EOF" in content[-1024:], "PDF não foi finalizado corretamente"
    assert len(content) > 1000


def test_summary_pdf_handles_empty_period():
    empty = ReportData(period=ReportPeriod.last_days(7), alerts=[], actions=[], audit_entries=[])
    content = build_summary_pdf(empty)
    assert content.startswith(b"%PDF-")


# -- janelas de tempo -----------------------------------------------------


def test_month_period_spans_exactly_one_month():
    period = ReportPeriod.month(2026, 3)
    assert period.start == datetime(2026, 3, 1, tzinfo=timezone.utc)
    assert period.end == datetime(2026, 4, 1, tzinfo=timezone.utc)


def test_december_period_rolls_over_the_year():
    """Caso de borda clássico: dezembro precisa virar janeiro do ano seguinte."""
    period = ReportPeriod.month(2026, 12)
    assert period.start == datetime(2026, 12, 1, tzinfo=timezone.utc)
    assert period.end == datetime(2027, 1, 1, tzinfo=timezone.utc)


def test_last_days_window_is_the_requested_length():
    period = ReportPeriod.last_days(15)
    assert (period.end - period.start).days == 15
